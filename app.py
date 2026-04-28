import streamlit as st
from docx import Document
from docx2pdf import convert
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import os
import platform
import pythoncom

st.set_page_config(page_title="Resume Generator", page_icon="📄")
st.title("Resume ATS Optimizer")

script_dir = os.path.dirname(os.path.abspath(__file__))
template_path = os.path.join(script_dir, "Template.docx")


def clear_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = ""


def set_run_style(run, size=8, bold=False, font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold


def replace_simple(paragraph, old_text, new_text, size=8, bold=False, align=None):
    clear_paragraph(paragraph)
    run = paragraph.add_run(new_text)
    set_run_style(run, size=size, bold=bold)

    if align is not None:
        paragraph.alignment = align


def replace_multiline(paragraph, text, size=8, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    """
    يستبدل الـ placeholder بعدة أسطر بنفس تنسيق مضبوط
    بدون ما يرث تنسيق العنوان أو الـ placeholder
    """
    lines = [line.strip() for line in str(text).splitlines() if line.strip()]

    parent = paragraph._element.getparent()
    index = parent.index(paragraph._element)

    # أول سطر مكان الـ placeholder
    clear_paragraph(paragraph)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1

    if lines:
        run = paragraph.add_run(lines[0])
        set_run_style(run, size=size, bold=bold)

    # باقي الأسطر
    for line in lines[1:]:
        new_p = paragraph.insert_paragraph_before("")
        parent.remove(new_p._element)
        parent.insert(index + 1, new_p._element)
        index += 1

        new_p.alignment = align
        new_p.paragraph_format.space_before = Pt(0)
        new_p.paragraph_format.space_after = Pt(0)
        new_p.paragraph_format.line_spacing = 1

        run = new_p.add_run(line)
        set_run_style(run, size=size, bold=bold)


def replace_text_in_paragraph(paragraph, mapping):
    full_text = ''.join(run.text for run in paragraph.runs)

    for key, config in mapping.items():
        if key in full_text:
            value = str(config["value"])
            size = config.get("size", 8)
            bold = config.get("bold", False)
            align = config.get("align", WD_ALIGN_PARAGRAPH.LEFT)

            if "\n" in value:
                replace_multiline(
                    paragraph=paragraph,
                    text=value,
                    size=size,
                    bold=bold,
                    align=align
                )
            else:
                new_text = full_text.replace(key, value)
                replace_simple(
                    paragraph=paragraph,
                    old_text=full_text,
                    new_text=new_text,
                    size=size,
                    bold=bold,
                    align=align
                )
            break


def replace_text_in_doc(doc, mapping):
    for paragraph in doc.paragraphs:
        replace_text_in_paragraph(paragraph, mapping)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_paragraph(paragraph, mapping)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            replace_text_in_paragraph(paragraph, mapping)
        for paragraph in section.footer.paragraphs:
            replace_text_in_paragraph(paragraph, mapping)


job_title = st.text_input(
    "Job Title",
    "Data Scientist | Data Engineer"
)

summary = st.text_area(
    "Summary",
    """Data Science Honors Graduate with a strong foundation in the end-to-end data lifecycle, from building scalable data infrastructure to developing predictive machine learning models. Proficient in designing Medallion architectures and applying statistical analysis to support data-driven decision-making. Skilled in Python, SQL, and Power BI, with the ability to translate complex datasets into actionable insights."""
)

tetco_bullets = st.text_area(
    "TETCO Bullets",
    """• Performed exploratory data analysis (EDA) to identify trends and support business insights.
• Architected automated ETL pipelines (SSIS) and 3-tier Medallion Architectures to ensure a Single Source of Truth and high-performance analytics structures.
• Optimized complex SQL (CTEs, Window Functions) and data quality protocols, reducing latency and ensuring 100% data reliability.
• Designed interactive Power BI dashboards with Figma UI/UX to simplify complex data navigation and translate findings into clear Data Stories for executives.
• Identified key data patterns and correlations during ingestion to provide evidence-based recommendations for operational improvements.
• Collaborated with cross-functional teams to define KPIs and documented technical Data Lineage to ensure infrastructure sustainability."""
)

tuwaiq_bullets = st.text_area(
    "Tuwaiq Bullets",
    """• Hands-on training on GCP services including BigQuery, Dataflow, and Cloud Storage.
• Building cloud-native ETL pipelines and scalable data architectures.
• Working with real-world datasets to support data processing and analytics."""
)

skills = st.text_area(
    "Skills",
    """• Data Science & Modeling: Machine Learning, Statistical Modeling, Predictive Analytics, Model Validation, Feature Engineering, Hypothesis Testing.
• Data Engineering & Architecture: ETL/ELT Pipelines, Medallion Architecture (Bronze/Silver/Gold), Data Modeling (Star/Snowflake).
• Programming & Databases: Advanced SQL (CTEs, Window Functions, T-SQL), Python (Pandas, Scikit-learn, NumPy), Java, Database Design.
• Tools & Infrastructure: Apache Airflow, Docker, AWS S3, SSIS, Jupyter Notebooks, Databricks Basic, Git/GitHub.
• Analytics & Business Intelligence: Power BI (DAX), Exploratory Data Analysis (EDA), KPI Identification, Data Storytelling, Figma UI/UX.
• Soft Skills: Analytical Thinking, Problem-Solving, Effective Communication, Professional Curiosity & Continuous Learning."""
)

proj1 = st.text_area(
    "Project 1",
    """• Orchestrated a complete ETL pipeline using a Medallion Architecture (Bronze, Silver, Gold) to ingest and refine complex student attendance and occupancy datasets.
• Developed optimized data schemas and Star/Snowflake modeling to enforce business logic, ensuring 100% data reliability for executive reporting.
• Engineered an interactive Power BI dashboard with Figma UI/UX to visualize spatial distribution and density trends, enabling data-driven classroom management."""
)

proj2 = st.text_area(
    "Project 2",
    """• Developed a hybrid IDS using XGBoost for known attacks (DoS, DDoS) and Variational Autoencoder for anomaly and zero-day attack detection.
• Performed feature engineering and evaluated models using Accuracy, Precision, Recall, and F1-score."""
)


if st.button("توليد السيرة الذاتية"):
    if not os.path.exists(template_path):
        st.error("ملف Template.docx غير موجود في نفس مجلد الكود.")
    else:
        try:
            doc = Document(template_path)

            data = {
                "{{JOB_TITLE}}": {
                    "value": job_title,
                    "size": 10,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.CENTER
                },
                "{{SUMMARY}}": {
                    "value": summary,
                    "size": 7,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.LEFT
                },
                "{{TETCO_BulletPoints}}": {
                    "value": tetco_bullets,
                    "size": 7,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.LEFT
                },
                "{{Tuwaiq_BulletPoints}}": {
                    "value": tuwaiq_bullets,
                    "size": 7,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.LEFT
                },
                "{{SKILLS_LIST}}": {
                    "value": skills,
                    "size": 7,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.LEFT
                },
                "{{PROJECT_1_CONTENT}}": {
                    "value": proj1,
                    "size": 7,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.LEFT
                },
                "{{PROJECT_2_CONTENT}}": {
                    "value": proj2,
                    "size": 7,
                    "bold": False,
                    "align": WD_ALIGN_PARAGRAPH.LEFT
                }
            }

            replace_text_in_doc(doc, data)

            output_docx = os.path.join(script_dir, "Final_Resume.docx")
            output_pdf = os.path.join(script_dir, "Final_Resume.pdf")

            doc.save(output_docx)

            with open(output_docx, "rb") as file:
                st.download_button(
                    label="تحميل Word",
                    data=file,
                    file_name="Final_Resume.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

            if platform.system() == "Windows":
                pythoncom.CoInitialize()
                convert(output_docx, output_pdf)

                with open(output_pdf, "rb") as file:
                    st.download_button(
                        label="تحميل PDF",
                        data=file,
                        file_name="Final_Resume.pdf",
                        mime="application/pdf"
                    )

                st.success("تم إنشاء Word و PDF بنجاح.")
            else:
                st.warning("تم إنشاء Word فقط. تحويل PDF يحتاج Windows + Microsoft Word.")

        except Exception as e:
            st.error(f"حدث خطأ: {e}")